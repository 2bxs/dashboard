import streamlit as st
import pandas as pd
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx2pdf import convert
import os
import zipfile
import io
import datetime

# Configuración de la página
st.set_page_config(page_title="Generador de Notificaciones de Reexpedición", layout="wide")
st.title("Generador Masivo de Notificaciones (Reexpediciones)")

# --- FUNCION REUTILIZABLE PARA GENERAR EL DOCUMENTO ---
def procesar_documento(rut, nrm, fecha_gen, gene, razon_social, df_completo, template_file):
    grupo_rut = df_completo[df_completo['RUT'] == rut]
    
    correo = grupo_rut['CORREO_FINAL'].iloc[0] if 'CORREO_FINAL' in grupo_rut.columns else ""
    tele = grupo_rut['TELE_FINAL'].iloc[0] if 'TELE_FINAL' in grupo_rut.columns else ""

    template_file.seek(0)
    doc = Document(template_file)
    
    replacements = {
        "[NOM_IMPOR]": str(razon_social),
        "[RUT_IMP]": str(rut),
        "[TELE]": str(tele),
        "[CORREO]": str(correo),
        "[NRM]": str(nrm),
        "[FECHA_GEN]": str(fecha_gen),
        "[FECH_GEN]": str(fecha_gen), 
        "[GENE]": str(gene)
    }
    
    # 1. Reemplazar en párrafos normales
    for p in doc.paragraphs:
        for key, val in replacements.items():
            if key in p.text:
                p.text = p.text.replace(key, val)
                for run in p.runs:
                    run.font.name = 'Tahoma'
                    run.font.size = Pt(9)
                    run.font.bold = True
                
    # 2. Operaciones dentro de las tablas
    if doc.tables: 
        for table in doc.tables:
            for row in table.rows:
                if any("[COD_RE]" in cell.text for cell in row.cells):
                    row._element.getparent().remove(row._element) 

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for key, val in replacements.items():
                            if key in p.text:
                                p.text = p.text.replace(key, val)
                                for run in p.runs:
                                    run.font.name = 'Tahoma'
                                    run.font.size = Pt(9)
                                    run.font.bold = True

        def formatear_fecha(valor):
            if pd.isna(valor) or str(valor).strip() in ["", "NaT"]:
                return ""
            try:
                return pd.to_datetime(valor).strftime("%d-%m-%Y")
            except:
                return str(valor).split(" ")[0]

        tabla_dinamica = doc.tables[0]
        for _, reexp_row in grupo_rut.iterrows():
            cells = tabla_dinamica.add_row().cells
            
            if len(cells) >= 4: 
                cells[0].text = str(reexp_row.get('REEXP.', ''))
                cells[1].text = formatear_fecha(reexp_row.get('Fecha Documento/ Visación'))
                cells[2].text = formatear_fecha(reexp_row.get('Fecha Control Salida'))
                cells[3].text = str(reexp_row.get('Avanzada Aduana', ''))
            
            if len(cells) >= 6:
                cells[4].text = str(nrm)
                cells[5].text = str(fecha_gen)

            for cell in cells:
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER 
                    for run in paragraph.runs:
                        run.font.name = 'Tahoma' 
                        run.font.size = Pt(9)   
                        run.font.bold = True    

    os.makedirs("temp_docs", exist_ok=True)
    docx_path = f"temp_docs/{nrm}.docx"
    pdf_path = f"temp_docs/{nrm}.pdf"
    doc.save(docx_path)
    
    try:
        convert(docx_path, pdf_path)
        return pdf_path, "pdf"
    except Exception:
        return docx_path, "docx"

# ==========================================
# ESTRUCTURA VISUAL EN CUADRÍCULA (3 COLUMNAS)
# ==========================================
col1, col2, col3 = st.columns([1.2, 1, 1.2], gap="large")

with col1:
    st.subheader("📂 1. Carga de Archivos")
    excel_file = st.file_uploader("Archivo 1 (Excel de registros)", type=["xlsx"])
    template_file = st.file_uploader("Archivo 2 (Plantilla base en .docx)", type=["docx"])

if excel_file and template_file:
    # Procesamiento inicial de datos
    df = pd.read_excel(excel_file)
    df.columns = df.columns.str.strip()
    
    if 'FONO 1' in df.columns:
        df['TELE_FINAL'] = df['FONO 1']
        if 'FONO 2' in df.columns: df['TELE_FINAL'] = df['TELE_FINAL'].fillna(df['FONO 2'])
        if 'FONO 3' in df.columns: df['TELE_FINAL'] = df['TELE_FINAL'].fillna(df['FONO 3'])
    else:
        df['TELE_FINAL'] = "SIN FONO"

    if 'CORREO 1' in df.columns:
        df['CORREO_FINAL'] = df['CORREO 1']
        if 'CORREO 2' in df.columns: df['CORREO_FINAL'] = df['CORREO_FINAL'].fillna(df['CORREO 2'])
        if 'CORREO 3' in df.columns: df['CORREO_FINAL'] = df['CORREO_FINAL'].fillna(df['CORREO 3'])
    else:
        df['CORREO_FINAL'] = "SIN CORREO"

    # Corrección de la columna 'Razón Social' aplicada aquí
    resumen = df.groupby('RUT').agg(
        Razon_Social=('Razón Social', 'first'),
        Cantidad=('RUT', 'count')
    ).reset_index()
    resumen = resumen.rename(columns={'Razon_Social': 'Nombre de Empresa', 'Cantidad': 'Cantidad de Reexpediciones'})

    def actualizar_nrm():
        inicio = st.session_state.global_nrm
        if inicio > 0:
            for i, r in enumerate(resumen['RUT']):
                st.session_state[f"nrm_{r}"] = f"MR-{(inicio + i):03d}"
        else:
            for r in resumen['RUT']:
                st.session_state[f"nrm_{r}"] = ""

    # Precarga de fechas en memoria
    meses = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"]
    hoy = datetime.datetime.now()
    fecha_actual_texto = f"{hoy.day:02d} de {meses[hoy.month - 1]} de {hoy.year}"
    
    for i, row in resumen.iterrows():
        rut = row['RUT']
        if f"fec_{rut}" not in st.session_state:
            st.session_state[f"fec_{rut}"] = fecha_actual_texto

    # Llenado de la Columna 2
    with col2:
        st.subheader("⚙️ 2. Configuración")
        st.number_input("Inicio de MR (NRM Inicial)", min_value=0, value=0, step=1, key="global_nrm", on_change=actualizar_nrm)
        global_gene = st.text_input("Iniciales de quien genera", key="global_iniciales")

    # Llenado de la Columna 3
    with col3:
        st.subheader("📊 3. Resumen y Generación")
        
        # Cuadros de métricas
        met1, met2 = st.columns(2)
        met1.metric("Empresas Totales", len(resumen))
        met2.metric("Reexp. Totales", int(resumen['Cantidad de Reexpediciones'].sum()))
        
        # Botón masivo trasladado arriba
        st.write("") # Espaciador
        if st.button("🚀 Generar Todo Masivamente", type="primary", use_container_width=True):
            with st.spinner("Generando documentos masivos..."):
                generated_files = []
                
                # Se lee directamente de la memoria para generar sin depender de la tabla visual
                for index, row in resumen.iterrows():
                    rut = row['RUT']
                    nrm_actual = st.session_state.get(f"nrm_{rut}", "")
                    fecha_actual = st.session_state.get(f"fec_{rut}", fecha_actual_texto)
                    
                    if nrm_actual: 
                        file_path, _ = procesar_documento(
                            rut, nrm_actual, fecha_actual, global_gene, row['Nombre de Empresa'], df, template_file
                        )
                        generated_files.append(file_path)
                
                if generated_files:
                    zip_buffer = io.BytesIO()
                    with zipfile.ZipFile(zip_buffer, "w") as zip_file:
                        for file_path in generated_files:
                            zip_file.write(file_path, os.path.basename(file_path))
                    
                    st.success("¡Éxito!")
                    st.download_button(
                        label="📥 Descargar ZIP Completo",
                        data=zip_buffer.getvalue(),
                        file_name="Notificaciones_Generadas.zip",
                        mime="application/zip",
                        use_container_width=True
                    )

    st.divider()
    st.subheader("📋 Registros Pendientes")
    
    # ==========================================
    # TABLA CON BORDE PLOMO (Contenedor)
    # ==========================================
    with st.container(border=True):
        col_widths = [2.2, 1.2, 0.6, 1.2, 1.5, 1.0, 1.0] 
        
        h1, h2, h3, h4, h5, h6, h7 = st.columns(col_widths)
        h1.markdown("**Empresa**")
        h2.markdown("**RUT**")
        h3.markdown("**Cant.**")
        h4.markdown("**NRM**")
        h5.markdown("**Fecha Gen.**")
        h6.markdown("**Acción**")
        h7.markdown("**Archivo**")
        
        st.markdown("---") 

        for index, row in resumen.iterrows():
            c1, c2, c3, c4, c5, c6, c7 = st.columns(col_widths)
            rut = row['RUT']
            
            c1.write(row['Nombre de Empresa'])
            c2.write(rut)
            c3.write(str(row['Cantidad de Reexpediciones']))
            
            nrm = c4.text_input("NRM", key=f"nrm_{rut}", label_visibility="collapsed")
            fecha_gen = c5.text_input("Fecha", key=f"fec_{rut}", label_visibility="collapsed")
            
            key_estado = f"file_data_{rut}"
            
            with c6:
                if st.button("⚙️ Generar", key=f"btn_gen_{rut}"):
                    if nrm: 
                        with st.spinner("⏳"):
                            file_path, ext = procesar_documento(rut, nrm, fecha_gen, global_gene, row['Nombre de Empresa'], df, template_file)
                            
                            with open(file_path, "rb") as f:
                                st.session_state[key_estado] = {
                                    "bytes": f.read(),
                                    "name": os.path.basename(file_path),
                                    "mime": "application/pdf" if ext == "pdf" else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                                }
            
            with c7:
                if key_estado in st.session_state:
                    file_info = st.session_state[key_estado]
                    st.download_button(
                        label="📥 PDF",
                        data=file_info["bytes"],
                        file_name=file_info["name"],
                        mime=file_info["mime"],
                        key=f"btn_dl_{rut}"
                    )
