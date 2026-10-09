import streamlit as st
import pandas as pd
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx2pdf import convert
import os
import zipfile
import io
import datetime

# Configuración de la página
st.set_page_config(page_title="Generador de Notificaciones de Reexpedición", layout="wide")

# Estilos visuales personalizados
st.markdown("""
<style>
/* Botones alineados a la izquierda, mínimo mitad de ancho pero expansibles si el texto es largo */
div.stButton > button, div.stDownloadButton > button {
    min-width: 50% !important;
    width: max-content !important;
    display: block !important;
    margin-right: auto !important;
}
/* Métricas con bordes redondeados y color naranjo candy */
[data-testid="metric-container"] {
    border: 2px solid #FF8C00 !important;
    border-radius: 15px !important;
    padding: 15px !important;
}
/* Separadores visuales naranjo candy */
hr {
    border-bottom: 2px solid #FF8C00 !important;
}
</style>
""", unsafe_allow_html=True)

st.title("Generador Masivo de Notificaciones (Reexpediciones)")

# --- FUNCION REUTILIZABLE PARA GENERAR EL DOCUMENTO ---
def procesar_documento(rut, nrm, fecha_gen, gene, razon_social, df_completo, template_file):
    grupo_rut = df_completo[df_completo['RUT'] == rut]
    
    correo = grupo_rut['CORREO_FINAL'].iloc[0] if 'CORREO_FINAL' in grupo_rut.columns else ""
    tele = grupo_rut['TELE_FINAL'].iloc[0] if 'TELE_FINAL' in grupo_rut.columns else ""

    template_file.seek(0)
    doc = Document(template_file)
    
    replacements = {
        "[NOM_IMPOR]": str(razon_social),
        "[RUT_IMP]": str(rut),
        "[TELE]": str(tele),
        "[CORREO]": str(correo),
        "[NRM]": str(nrm),
        "[FECHA_GEN]": str(fecha_gen),
        "[FECH_GEN]": str(fecha_gen), 
        "[GENE]": str(gene)
    }
    
    for p in doc.paragraphs:
        for key, val in replacements.items():
            if key in p.text:
                p.text = p.text.replace(key, val)
                for run in p.runs:
                    run.font.name = 'Tahoma'
                    run.font.size = Pt(9)
                    run.font.bold = True
                
    if doc.tables: 
        for table in doc.tables:
            for row in table.rows:
                if any("[COD_RE]" in cell.text for cell in row.cells):
                    row._element.getparent().remove(row._element) 

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for key, val in replacements.items():
                            if key in p.text:
                                p.text = p.text.replace(key, val)
                                for run in p.runs:
                                    run.font.name = 'Tahoma'
                                    run.font.size = Pt(9)
                                    run.font.bold = True

        def formatear_fecha(valor):
            if pd.isna(valor) or str(valor).strip() in ["", "NaT"]:
                return ""
            try:
                return pd.to_datetime(valor).strftime("%d-%m-%Y")
            except:
                return str(valor).split(" ")[0]

        tabla_dinamica = doc.tables[0]
        for _, reexp_row in grupo_rut.iterrows():
            cells = tabla_dinamica.add_row().cells
            
            if len(cells) >= 4: 
                cells[0].text = str(reexp_row.get('REEXP.', ''))
                cells[1].text = formatear_fecha(reexp_row.get('Fecha Documento/ Visación'))
                cells[2].text = formatear_fecha(reexp_row.get('Fecha Control Salida'))
                cells[3].text = str(reexp_row.get('Avanzada Aduana', ''))
            
            if len(cells) >= 6:
                cells[4].text = str(nrm)
                cells[5].text = str(fecha_gen)

            for cell in cells:
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER 
                    for run in paragraph.runs:
                        run.font.name = 'Tahoma' 
                        run.font.size = Pt(9)   
                        run.font.bold = True    

    os.makedirs("temp_docs", exist_ok=True)
    docx_path = f"temp_docs/{nrm}.docx"
    pdf_path = f"temp_docs/{nrm}.pdf"
    doc.save(docx_path)
    
    try:
        convert(docx_path, pdf_path)
        return pdf_path, "pdf"
    except Exception:
        return docx_path, "docx"

# ==========================================
# 1. CARGA DE ARCHIVOS
# ==========================================
st.subheader("📂 1. Carga de Archivos")
col1, col2 = st.columns(2)
with col1:
    excel_file = st.file_uploader("Archivo 1 (Excel de registros)", type=["xlsx"])
with col2:
    template_file = st.file_uploader("Archivo 2 (Plantilla base en .docx)", type=["docx"])

if excel_file and template_file:
    df = pd.read_excel(excel_file)
    df.columns = df.columns.str.strip()
    
    if 'FONO 1' in df.columns:
        df['TELE_FINAL'] = df['FONO 1']
        if 'FONO 2' in df.columns: df['TELE_FINAL'] = df['TELE_FINAL'].fillna(df['FONO 2'])
        if 'FONO 3' in df.columns: df['TELE_FINAL'] = df['TELE_FINAL'].fillna(df['FONO 3'])
    else:
        df['TELE_FINAL'] = "SIN FONO"

    if 'CORREO 1' in df.columns:
        df['CORREO_FINAL'] = df['CORREO 1']
        if 'CORREO 2' in df.columns: df['CORREO_FINAL'] = df['CORREO_FINAL'].fillna(df['CORREO 2'])
        if 'CORREO 3' in df.columns: df['CORREO_FINAL'] = df['CORREO_FINAL'].fillna(df['CORREO 3'])
    else:
        df['CORREO_FINAL'] = "SIN CORREO"

    resumen = df.groupby('RUT').agg(
        Razon_Social=('Razón Social', 'first'),
        Cantidad=('RUT', 'count')
    ).reset_index()
    resumen = resumen.rename(columns={'Razon_Social': 'Nombre de Empresa', 'Cantidad': 'Cantidad de Reexpediciones'})

    # Funciones de actualización
    def actualizar_nrm():
        inicio = st.session_state.global_nrm
        if inicio > 0:
            for i, r in enumerate(resumen['RUT']):
                st.session_state[f"nrm_{r}"] = f"MR-{(inicio + i):03d}"
        else:
            for r in resumen['RUT']:
                st.session_state[f"nrm_{r}"] = ""

    meses = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"]
